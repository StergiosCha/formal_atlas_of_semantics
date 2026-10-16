#!/usr/bin/env python3
"""Export the two-page review brief and a bounded, hash-pinned source packet.

Uses the report's document tools. Never changes proofs, audits or the report.
Requires the already-rendered checkpoint-7 report PDF and Word document.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from render_report import STEM as REPORT_STEM, field

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
STEM = "Montague_Review_Brief_for_Robin_and_Rasmus_Checkpoint_7_2026-10-08"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--report-dir", type=Path, required=True)
    args = parser.parse_args()
    destination = args.output_dir.resolve()
    if destination.exists():
        raise RuntimeError("Use a new output directory to preserve existing review snapshots.")
    report_dir = args.report_dir.resolve()
    if (report_dir / (REPORT_STEM + ".md")).read_bytes() != (HERE / "TECHNICAL_REPORT.md").read_bytes():
        raise RuntimeError("The exported technical report is not the current checkpoint report.")
    sources = []
    for branch in ("montague_core", "montague_legacy"):
        directory = ROOT / "deep" / branch
        audit = json.loads((directory / "verification.json").read_text())
        for name, data in audit["files"].items():
            if digest(directory / name) != data["sha256"]:
                raise RuntimeError(f"Audited proof changed: {branch}/{name}")
        sources.extend(p for p in directory.iterdir()
                       if p.is_file() and (p.suffix in {".v", ".md", ".py", ".json"}
                                          or p.name == "_CoqProject"))
    legacy_audit = json.loads((ROOT / "deep/montague_legacy/verification.json").read_text())
    if digest(ROOT / "deep/PTQ_deep2.v") != legacy_audit["legacy_source"]["sha256"]:
        raise RuntimeError("The legacy source differs from its audited pin.")
    sources.extend([ROOT / "deep/PTQ_deep2.v", ROOT / "atlas_data/verify.py"])
    destination.mkdir(parents=True)
    for source in sorted(sources):
        target = destination / "proofs" / source.relative_to(ROOT)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
    for suffix in (".pdf", ".docx"):
        shutil.copy2(report_dir / (REPORT_STEM + suffix), destination / (REPORT_STEM + suffix))

    markdown = (HERE / "REVIEW_BRIEF.md").read_text()
    if "\u2014" in markdown or "\u2013" in markdown:
        raise RuntimeError("The review brief must not contain em or en dashes.")

    def link(match):
        label, location = match.groups()
        source = (HERE / location).resolve()
        if source == HERE / "TECHNICAL_REPORT.md":
            target = REPORT_STEM + ".pdf"
        else:
            target = "proofs/" + source.relative_to(ROOT).as_posix()
        if not (destination / target).is_file():
            raise RuntimeError(f"Missing packaged link: {target}")
        return f"[{label}]({target})"

    packaged = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", link, markdown)
    (destination / (STEM + ".md")).write_text(packaged)
    word = destination / (STEM + ".docx")
    subprocess.run(["pandoc", "--standalone", "--from=markdown", "--to=docx", "-o", str(word)],
                   input=packaged, text=True, check=True)
    doc = Document(word)
    doc.core_properties.title = "Montague in the Formal Atlas: review brief"
    doc.core_properties.subject = "Checkpoint 7 scientific review; no reviewer endorsement claimed"
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.62)
        section.left_margin = section.right_margin = Inches(0.72)
        section.header_distance = section.footer_distance = Inches(0.25)
        header = section.header.paragraphs[0]
        header.text = "FORMAL ATLAS | REVIEW COPY | CHECKPOINT 7"
        for run in header.runs:
            run.font.name, run.font.size = "Arial", Pt(8)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Review brief | ")
        field(footer, "PAGE")
        footer.add_run(" / ")
        field(footer, "NUMPAGES")
        for run in footer.runs:
            run.font.size = Pt(8)
    styles = {style.style_id: style for style in doc.styles}
    normal = styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.line_spacing = 1.02
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for level, size in ((1, 17), (2, 11.5)):
        style = styles[f"Heading{level}"]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string("243B53")
        style.paragraph_format.space_before = Pt(8 if level == 2 else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    for name in ("SourceCode", "VerbatimChar"):
        if name in styles:
            styles[name].font.name, styles[name].font.size = "Consolas", Pt(8.5)
    for paragraph in doc.paragraphs:
        if paragraph.text == "Assumptions and review decision":
            paragraph.paragraph_format.page_break_before = True
    doc.save(word)
    with tempfile.TemporaryDirectory(prefix="montague-brief-office-") as profile:
        subprocess.run(["soffice", "-env:UserInstallation=" + Path(profile).as_uri(), "--headless",
                        "--convert-to", "pdf", "--outdir", str(destination), str(word)],
                       check=True, timeout=90)
    pdf = word.with_suffix(".pdf")
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    if not re.search(r"(?m)^Pages:\s+2\s*$", info):
        raise RuntimeError("The review brief is not exactly two pages; review its layout.")
    extracted = subprocess.check_output(["pdftotext", str(pdf), "-"], text=True)
    if "\ufffd" in extracted or "Three questions for your review" not in extracted:
        raise RuntimeError("PDF text validation failed.")
    manifest = {
        "checkpoint": 7,
        "status": "frozen review snapshot, not a published or approved release",
        "reviewer_approval": False,
        "git_head_for_context_only": subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "identity_note": "Source includes uncommitted work; use these hashes, not Git HEAD, to identify it.",
        "files": {p.relative_to(destination).as_posix(): digest(p)
                  for p in sorted(destination.rglob("*")) if p.is_file()},
    }
    (destination / "MANIFEST.json").write_text(json.dumps(manifest, indent=2) + "\n")
    archive = destination.with_suffix(".zip")
    if archive.exists():
        raise RuntimeError("Refusing to overwrite an existing review archive.")
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as bundle:
        for p in sorted(destination.rglob("*")):
            if p.is_file():
                bundle.write(p, Path(destination.name) / p.relative_to(destination))
    print(info)
    print(word)
    print(pdf)
    print(archive)


if __name__ == "__main__":
    main()
