#!/usr/bin/env python3
"""Build the code-led walkthrough from a frozen review packet, not edited proofs.

Requires pandoc, python-docx, LibreOffice, poppler and Coq 8.20.1.
Validates source hashes, checks the copied library and captures the demo.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from render_report import field

HERE = Path(__file__).resolve().parent
STEM = "Montague_Code_Walkthrough_for_Robin_and_Rasmus_2026-10-09"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def checked(command, **kwargs):
    result = subprocess.run(command, text=True, stdout=subprocess.PIPE,
                            stderr=subprocess.STDOUT, timeout=360, **kwargs)
    if result.returncode or re.search(r"\bError:", result.stdout):
        raise RuntimeError(f"Command failed: {command}\n{result.stdout}")
    return result.stdout


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot-dir", type=Path, required=True, help="the frozen packet's proofs directory")
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    snapshot, destination = args.snapshot_dir.resolve(), args.output_dir.resolve()
    if destination == snapshot or snapshot in destination.parents:
        raise RuntimeError("Never write the walkthrough into its frozen input snapshot.")
    if (destination / "MANIFEST.json").exists():
        raise RuntimeError("Use a new output directory for a new completed document edition.")
    original_manifest = json.loads((snapshot.parent / "MANIFEST.json").read_text())
    for name, expected in original_manifest["files"].items():
        if name.startswith("proofs/") and digest(snapshot.parent / name) != expected:
            raise RuntimeError(f"Frozen input changed: {name}")
    sys.path.insert(0, str(snapshot / "atlas_data"))
    import verify
    audits, sources = {}, {}
    for branch in ("montague_core", "montague_legacy"):
        audit = json.loads((snapshot / "deep" / branch / "verification.json").read_text())
        audits[branch] = audit
        for name, data in audit["files"].items():
            relative = f"deep/{branch}/{name}"
            if digest(snapshot / relative) != data["sha256"]:
                raise RuntimeError(f"Source no longer matches its proof audit: {relative}")
            declarations, warnings = verify.parse(snapshot / relative)
            if warnings:
                raise RuntimeError(f"Unparsed source: {relative}")
            sources[relative] = {
                "text": (snapshot / relative).read_text(),
                "declarations": declarations,
                "audit": data,
                "branch": branch,
            }
    assert len(sources) == 26
    assert sum(a['proved_statement_count'] for a in audits.values()) == 249
    template = (HERE / "WALKTHROUGH.md").read_text()
    if "\u2014" in template or "\u2013" in template:
        raise RuntimeError("Use no em or en dashes in the newly authored prose.")
    excerpts = []

    def snippet(match):
        relative, start, end = match.groups()
        item = sources[relative]
        declarations = item["declarations"]
        starts = [i for i, d in enumerate(declarations) if d['name'] == start]
        ends = [i for i, d in enumerate(declarations) if d['name'] == end]
        if len(starts) != 1 or len(ends) != 1 or ends[0] < starts[0]:
            raise RuntimeError(f"Ambiguous or reversed source excerpt: {match.group(0)}")
        first = declarations[starts[0]]["line"]
        last = (declarations[ends[0]+1]['line'] - 1
                if ends[0]+1 < len(declarations) else len(item['text'].splitlines()))
        code = '\n'.join(item['text'].splitlines()[first-1:last]).rstrip()
        excerpts.append({"file": relative, "first_line": first, "last_line": last,
                         "from": start, "through": end, "source_sha256": item['audit']['sha256']})
        return (f"Source: [{relative}](proofs/{relative}), lines {first}-{last}.\n\n"
                f"```coq\n{code}\n```\n")

    body = re.sub(r"\{\{snippet (\S+) (\w+) (\w+)\}\}", snippet, template)
    session = (HERE / "WALKTHROUGH_SESSION.txt").read_text()
    body = body.replace("{{session}}", f"```coq\n{session.rstrip()}\n```\n")
    appendix = []
    for n, (relative, item) in enumerate(sources.items(), 1):
        appendix.extend([f"## A.{n}. {relative}", "",
                         f"Source SHA256: `{item['audit']['sha256']}`.", "",
                         "Named proved statements (explicit hypotheses remain in their types):", ""])
        for name in item['audit']['proved_statements']:
            if item['branch'] == 'montague_core':
                assumption = "closed global context"
            else:
                dependencies = audits['montague_legacy']['global_dependencies_by_statement'][name]
                assumption = (f"{len(dependencies)} retained legacy global parameters"
                              if dependencies else "closed global context")
            appendix.append(f"- `{name}`: {assumption}.")
        appendix.extend(["", f"```coq\n{item['text'].rstrip()}\n```", ""])
    body = body.replace("{{all_sources}}", '\n'.join(appendix))
    if '{{' in body:
        raise RuntimeError("Unexpanded document directive")
    destination.mkdir(parents=True, exist_ok=True)
    copied = destination / "proofs"
    shutil.copytree(snapshot, copied, dirs_exist_ok=True,
                    ignore=shutil.ignore_patterns('__pycache__', '*.pyc', '*.vo', '*.vos', '*.vok', '*.glob', '.*.aux'))
    (destination / (STEM + '.md')).write_text(body)
    (destination / 'EXCERPTS.json').write_text(json.dumps(excerpts, indent=2) + '\n')
    for name in ('WALKTHROUGH.md', 'WALKTHROUGH_SESSION.txt', 'render_walkthrough.py', 'render_report.py'):
        target = destination / "renderer" / name
        target.parent.mkdir(exist_ok=True)
        shutil.copy2(HERE / name, target)
    shutil.copy2(HERE / 'WALKTHROUGH_SESSION.txt', destination / 'WALKTHROUGH_SESSION.txt')

    print('Checking the copied core and capturing the companion inspection session.', flush=True)
    core_log = checked([sys.executable, 'deep/montague_core/check.py', '--check'], cwd=copied)
    (destination / 'CORE_CHECK.txt').write_text(core_log)
    print('Core: fresh build, kernel and negative checks passed.', flush=True)
    session_log = checked([sys.executable, 'deep/montague_legacy/check.py', '--check', '--inspect'],
                          cwd=copied, input=session)
    # Coq may print the unreduced wdomain ... wt alias instead of bool.
    # Check the four computed constructors, not the pretty-printer's type spelling.
    results = re.findall(r'=\s*(true|false)\s*\n\s*:', session_log)
    if results != ['true', 'false', 'true', 'false']:
        raise RuntimeError(f"Unexpected walkthrough computations: {results}\n{session_log}")
    (destination / 'WALKTHROUGH_SESSION_LOG.txt').write_text(session_log)
    print('Companion: kernel checks passed; four tutorial computations matched.', flush=True)

    word = destination / (STEM + '.docx')
    subprocess.run(['pandoc', str(destination / (STEM+'.md')), '--standalone', '--from=markdown',
                    '--to=docx', '--no-highlight', '-o', str(word)], check=True)
    doc = Document(word)
    styles = {s.style_id: s for s in doc.styles}
    normal = styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.line_spacing = 1.06
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for level, size in ((1, 18), (2, 13), (3, 11.5)):
        style = styles[f'Heading{level}']
        style.font.name, style.font.size = 'Calibri', Pt(size)
        style.font.color.rgb = RGBColor.from_string('243B53')
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    for name in ('SourceCode', 'VerbatimChar'):
        if name in styles:
            styles[name].font.name, styles[name].font.size = 'Menlo', Pt(8)
    if 'SourceCode' in styles:
        code = styles['SourceCode'].paragraph_format
        code.line_spacing = 1.0
        code.space_after = Pt(6)
        code.keep_together = False
        code.keep_with_next = False
        code.widow_control = False
    styles['Title'].font.name, styles['Title'].font.size = 'Calibri', Pt(29)
    styles['Title'].paragraph_format.space_before = Pt(80)
    styles['Subtitle'].font.size = Pt(17)
    headings = [p for p in doc.paragraphs if p.style.style_id == 'Heading1']
    first = headings[0]
    first.insert_paragraph_before(
        'A guided explanation of the existing proofs, with exact source excerpts, '
        'complete modules and a reproducible inspection session.\n'
        'Checkpoint 7 remains unchanged. No independent reviewer endorsement is claimed.')
    contents = first.insert_paragraph_before('Reading map', styles['Heading1'])
    contents.paragraph_format.page_break_before = True
    for p in headings:
        first.insert_paragraph_before(p.text)
        p.paragraph_format.page_break_before = True
    for p in doc.paragraphs:
        if p.style.style_id == 'Heading2' and p.text.startswith('A.'):
            p.paragraph_format.page_break_before = True
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(.68)
        section.left_margin = section.right_margin = Inches(.68)
        section.header_distance = section.footer_distance = Inches(.28)
        section.different_first_page_header_footer = True
        section.header.paragraphs[0].text = 'FORMAL ATLAS | CODE-LED WALKTHROUGH | CHECKPOINT 7'
        for run in section.header.paragraphs[0].runs:
            run.font.size = Pt(8)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('Review copy | ')
        field(footer, 'PAGE'); footer.add_run(' / '); field(footer, 'NUMPAGES')
        for run in footer.runs:
            run.font.size = Pt(8)
    for table in doc.tables:
        table.style = styles['Table']
        for n, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if n == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
    doc.save(word)
    with tempfile.TemporaryDirectory(prefix='montague-walkthrough-office-') as profile:
        print(checked(['soffice', '-env:UserInstallation='+Path(profile).as_uri(), '--headless',
                       '--convert-to', 'pdf', '--outdir', str(destination), str(word)]), flush=True)
    pdf = word.with_suffix('.pdf')
    info = checked(['pdfinfo', str(pdf)])
    text = checked(['pdftotext', '-layout', str(pdf), '-'])
    if '\ufffd' in text or '20. What a reviewer should decide' not in text:
        raise RuntimeError('PDF text validation failed')
    # Caches created by the checker are not distributable source or frozen inputs.
    files = [p for p in destination.rglob('*') if p.is_file() and '__pycache__' not in p.parts]
    manifest = {
        'document_date': '2026-10-09', 'mathematical_checkpoint': 7,
        'scope': 'Documentation of unchanged frozen source; not a new formalization or release',
        'source_snapshot_manifest_sha256': digest(snapshot.parent/'MANIFEST.json'),
        'source_files': len(sources), 'named_proved_statements': 249,
        'exact_source_excerpts': len(excerpts), 'negative_fixtures_rejected': 37,
        'interactive_boolean_results': results, 'reviewer_approval': False,
        'files': {p.relative_to(destination).as_posix(): digest(p) for p in sorted(files)},
    }
    (destination/'MANIFEST.json').write_text(json.dumps(manifest, indent=2)+'\n')
    archive = destination.with_suffix('.zip')
    if archive.exists():
        raise RuntimeError('Refusing to overwrite an existing walkthrough archive')
    with zipfile.ZipFile(archive, 'w', compression=zipfile.ZIP_DEFLATED) as bundle:
        for p in [*sorted(files), destination/'MANIFEST.json']:
            bundle.write(p, Path(destination.name)/p.relative_to(destination))
    print(info)
    print(word); print(pdf); print(archive)


if __name__ == '__main__':
    main()
