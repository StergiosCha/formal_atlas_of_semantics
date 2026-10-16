#!/usr/bin/env python3
"""Render the introductory two-pager, without modifying proofs or audits.

Requires pandoc, python-docx, LibreOffice and Poppler. Refuses to replace an
existing export directory. No source archive or long report is attached.
"""

import argparse
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "OVERVIEW_FOR_ROBIN_AND_RASMUS.md"
STEM = "Formal_Atlas_Overview_for_Robin_and_Rasmus_2026-10-16"
PAGE_TWO = "What is checked, and what needs discussion"


def page_field(paragraph, instruction):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    output = args.output_dir.resolve()
    if output.exists():
        raise RuntimeError("Choose a new export directory; existing snapshots are preserved.")
    output.mkdir(parents=True)
    markdown = SOURCE.read_text()
    if "\u2014" in markdown:
        raise RuntimeError("Use no em dashes in this document.")
    shutil.copy2(SOURCE, output / (STEM + ".md"))
    word = output / (STEM + ".docx")
    subprocess.run(
        ["pandoc", "--standalone", "--from=markdown", "--to=docx", "-o", str(word)],
        input=markdown, text=True, check=True,
    )
    doc = Document(word)
    doc.core_properties.title = "Formal Atlas of Semantics: a short overview"
    doc.core_properties.author = "Stergios Chatzikyriakidis"
    doc.core_properties.subject = "Discussion draft for Robin Cooper and Rasmus Blanck"
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.75)
        section.header_distance = section.footer_distance = Inches(0.28)
        header = section.header.paragraphs[0]
        header.text = "FORMAL ATLAS OF SEMANTICS | DISCUSSION COPY"
        for run in header.runs:
            run.font.name, run.font.size = "Calibri", Pt(8)
            run.font.color.rgb = RGBColor.from_string("526477")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("16 October 2026 | ")
        page_field(footer, "PAGE")
        footer.add_run(" / ")
        page_field(footer, "NUMPAGES")
        for run in footer.runs:
            run.font.name, run.font.size = "Calibri", Pt(8)
    styles = {style.style_id: style for style in doc.styles}
    for name in ("Normal", "BodyText", "FirstParagraph", "Compact"):
        if name not in styles:
            continue
        style = styles[name]
        style.font.name, style.font.size = "Calibri", Pt(10.5)
        style.paragraph_format.line_spacing = 1.02
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.widow_control = True
    for level, size in ((1, 18), (2, 11.5)):
        style = styles[f"Heading{level}"]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string("243B53")
        style.paragraph_format.space_before = Pt(8 if level == 2 else 0)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    if "VerbatimChar" in styles:
        styles["VerbatimChar"].font.name = "Consolas"
        styles["VerbatimChar"].font.size = Pt(9)
    for paragraph in doc.paragraphs:
        if paragraph.text == PAGE_TWO:
            paragraph.paragraph_format.page_break_before = True
    doc.save(word)
    with tempfile.TemporaryDirectory(prefix="atlas-overview-office-") as profile:
        subprocess.run(
            ["soffice", "-env:UserInstallation=" + Path(profile).as_uri(),
             "--headless", "--convert-to", "pdf", "--outdir", str(output), str(word)],
            check=True, timeout=60,
        )
    pdf = word.with_suffix(".pdf")
    info = subprocess.check_output(["pdfinfo", str(pdf)], text=True)
    pages = next(line.split(":", 1)[1].strip() for line in info.splitlines()
                 if line.startswith("Pages:"))
    if pages != "2":
        raise RuntimeError(f"Expected two pages, found {pages}; inspect layout.")
    text = subprocess.check_output(["pdftotext", str(pdf), "-"], text=True)
    if "\ufffd" in text or "Two questions for discussion" not in text:
        raise RuntimeError("PDF text validation failed.")
    metadata = {
        "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "date": "2026-10-16",
        "pages": 2,
        "status": "discussion draft, no reviewer endorsement or publication claimed",
        "files": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                  for p in sorted(output.iterdir()) if p.is_file()},
    }
    (output / "MANIFEST.json").write_text(json.dumps(metadata, indent=2) + "\n")
    print(f"Validated two-page PDF: {pdf}")
    print(f"Editable Word copy: {word}")


if __name__ == "__main__":
    main()
