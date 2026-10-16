#!/usr/bin/env python3
"""Export the reviewer report to Word and PDF. Requires pandoc, python-docx,
LibreOffice and pdfinfo. This is document generation, not a Coq check.
"""
import argparse
import json
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
STEM = "Montague_Proof_Theory_Report_for_Robin_and_Rasmus_Checkpoint_7_2026-10-08"


def field(paragraph, instruction):
    element = OxmlElement("w:fldSimple")
    element.set(qn("w:instr"), instruction)
    paragraph._p.append(element)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    destination = args.output_dir.resolve()
    destination.mkdir(parents=True, exist_ok=True)
    source = HERE / "TECHNICAL_REPORT.md"
    checkpoint = json.loads((HERE / "verification.json").read_text())
    legacy = HERE.parent / "montague_legacy"
    legacy_checkpoint = json.loads((legacy / "verification.json").read_text())
    total = checkpoint['proved_statement_count'] + legacy_checkpoint['proved_statement_count']
    word = destination / (STEM + ".docx")
    subprocess.run(["pandoc", str(source), "--standalone", "--from=markdown",
                    "--to=docx", "-o", str(word)], check=True)
    document = Document(word)
    styles = {style.style_id: style for style in document.styles}
    for section in document.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.78)
        section.left_margin = section.right_margin = Inches(0.85)
        section.header_distance = section.footer_distance = Inches(0.35)
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for level, size in ((1, 16), (2, 12), (3, 11)):
        style = styles[f"Heading{level}"]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("243B53")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(15 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)
    for name in ("SourceCode", "VerbatimChar"):
        if name in styles:
            styles[name].font.name = "Courier New"
            styles[name].font.size = Pt(8.5)
    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(27)
    title.font.color.rgb = RGBColor.from_string("243B53")
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(18)
    styles["Subtitle"].font.name = "Arial"
    styles["Subtitle"].font.size = Pt(15)
    styles["Subtitle"].paragraph_format.space_after = Pt(24)

    headings = [p for p in document.paragraphs if p.style.name == "Heading 1"]
    first = headings[0]
    first.paragraph_format.page_break_before = True
    note = first.insert_paragraph_before(
        f"Checkpoint 7 | {total} checked statements | "
        f"Coq {checkpoint['coq_version']}\n"
        f"{checkpoint['proved_statement_count']} closed-core statements; "
        f"{legacy_checkpoint['proved_statement_count']} separately audited companion statements.\n"
        "Guarded Prop/Boolean agreement; not unrestricted model equivalence.\n"
        "Conditional soundness; not full Montague or Henkin adequacy.\n"
        "Prepared for review. No reviewer endorsement is claimed.")
    note.paragraph_format.space_before = Pt(24)
    contents = first.insert_paragraph_before("Contents", styles["Heading1"])
    contents.paragraph_format.page_break_before = True
    for heading in headings:
        entry = first.insert_paragraph_before(heading.text)
        entry.paragraph_format.space_after = Pt(9)

    for table in document.tables:
        table.style = styles["Table"]
        for index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            if index == 0:
                properties.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                        if index == 0:
                            run.bold = True

    for section in document.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.text = "FORMAL ATLAS | Montague proof-theory workbench"
        header.style = styles["Normal"]
        for run in header.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Review copy | ")
        field(footer, "PAGE")
        footer.add_run(" / ")
        field(footer, "NUMPAGES")
        for run in footer.runs:
            run.font.size = Pt(8)
    document.save(word)
    shutil.copy2(source, destination / (STEM + ".md"))
    shutil.copy2(HERE / "verification.json", destination / "verification.json")
    shutil.copy2(HERE / "MODEL_SPEC.md", destination / "MODEL_SPEC.md")
    shutil.copy2(HERE / "INTENSIONALITY_SPEC.md", destination / "INTENSIONALITY_SPEC.md")
    shutil.copy2(HERE / "IL_SPEC.md", destination / "IL_SPEC.md")
    shutil.copy2(legacy / "LEGACY_SPEC.md", destination / "LEGACY_SPEC.md")
    shutil.copy2(legacy / "OBSERVATION_SPEC.md", destination / "OBSERVATION_SPEC.md")
    shutil.copy2(legacy / "OBSERVATION_SUITE.json", destination / "OBSERVATION_SUITE.json")
    shutil.copy2(legacy / "verification.json", destination / "legacy_verification.json")
    with tempfile.TemporaryDirectory(prefix="montague-report-office-") as profile:
        subprocess.run(["soffice", "-env:UserInstallation=" + Path(profile).as_uri(),
                        "--headless", "--convert-to", "pdf", "--outdir",
                        str(destination), str(word)], check=True, timeout=90)
    pdf = word.with_suffix(".pdf")
    if not pdf.is_file() or pdf.stat().st_size == 0:
        raise RuntimeError("LibreOffice did not produce the PDF")
    subprocess.run(["pdfinfo", str(pdf)], check=True)
    print(word)
    print(pdf)


if __name__ == "__main__":
    main()
